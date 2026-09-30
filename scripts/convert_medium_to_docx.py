"""
NexusLake Markdown to Word (.docx) Converter
Converts MEDIUM_TUTORIAL_BLOG_POST.md into a beautifully formatted Microsoft Word (.docx) document
with embedded images, styled tables, code callouts, custom typography, and headers.
"""

import os
import re
import sys

# Ensure UTF-8 output on Windows console
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex):
    """Sets background shading color for a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Sets internal padding for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="D1D5DB", sz="4"):
    """Applies subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_callout_borders(table, border_color="3B82F6"):
    """Sets a left-only border for callout blocks."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def set_code_block_borders(table, border_color="CBD5E1"):
    """Sets a subtle border for code blocks."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'  <w:insideH w:val="none"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_formatted_text(p, text, default_size=10.5, default_color=None):
    """
    Parses inline markdown tokens: **bold**, *italic*, `code`, and plain text.
    """
    # Regex to tokenize markdown elements
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)

    for token in tokens:
        if not token:
            continue

        run = p.add_run()
        run.font.name = "Calibri"
        run.font.size = Pt(default_size)
        if default_color:
            run.font.color.rgb = default_color
        else:
            run.font.color.rgb = RGBColor(31, 41, 55)  # #1F2937

        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(default_size - 1)
            run.font.color.rgb = RGBColor(185, 28, 28)  # Red accent for inline code
        elif token.startswith("[") and "](" in token and token.endswith(")"):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                run.text = m.group(1)
                run.font.color.rgb = RGBColor(37, 99, 235)  # Blue link
                run.underline = True
            else:
                run.text = token
        else:
            run.text = token


def convert_markdown_to_docx(md_path, docx_path, base_dir):
    print(f"[*] Reading Markdown file: {md_path}")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Set standard page margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Configure Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(31, 41, 55)

    i = 0
    total_lines = len(lines)

    while i < total_lines:
        raw_line = lines[i]
        line = raw_line.rstrip()

        # 1. Blank Line
        if not line:
            i += 1
            continue

        # 2. Horizontal Rule (---)
        if line == "---" or line == "***":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("―" * 45)
            run.font.color.rgb = RGBColor(209, 213, 219)
            i += 1
            continue

        # 3. Code Block (```)
        if line.startswith("```"):
            code_lang = line[3:].strip()
            code_lines = []
            i += 1
            while i < total_lines and not lines[i].rstrip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1  # Skip the closing ```

            code_text = "\n".join(code_lines)

            # Create a 1x1 table styled like a modern IDE code block
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = False
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.5)

            # Dark theme background for code blocks
            set_cell_background(cell, "0F172A")  # #0F172A dark slate
            set_code_block_borders(tbl, "334155")
            set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

            # Add language header label
            if code_lang:
                p_lang = cell.paragraphs[0]
                p_lang.paragraph_format.space_after = Pt(4)
                r_lang = p_lang.add_run(f"CODE [{code_lang.upper()}]")
                r_lang.font.name = "Consolas"
                r_lang.font.size = Pt(8.5)
                r_lang.bold = True
                r_lang.font.color.rgb = RGBColor(96, 165, 250)  # Light Blue

                p_code = cell.add_paragraph()
            else:
                p_code = cell.paragraphs[0]

            p_code.paragraph_format.space_before = Pt(0)
            p_code.paragraph_format.space_after = Pt(0)
            p_code.paragraph_format.line_spacing = 1.15
            r_code = p_code.add_run(code_text)
            r_code.font.name = "Consolas"
            r_code.font.size = Pt(9.0)
            r_code.font.color.rgb = RGBColor(241, 245, 249)  # Light Slate White

            # Add spacing after table
            p_post = doc.add_paragraph()
            p_post.paragraph_format.space_before = Pt(4)
            p_post.paragraph_format.space_after = Pt(4)
            continue

        # 4. Markdown Table (| ... |)
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < total_lines and lines[i].rstrip().startswith("|") and lines[i].rstrip().endswith("|"):
                table_lines.append(lines[i].rstrip())
                i += 1

            if len(table_lines) >= 2:
                # Parse rows and skip separator row (| :--- | :--- |)
                parsed_rows = []
                for t_idx, t_line in enumerate(table_lines):
                    # Check if separator row
                    cells = [c.strip() for c in t_line.strip("|").split("|")]
                    if t_idx == 1 and all(set(c).issubset({'-', ':', ' '}) for c in cells):
                        continue  # Skip divider
                    parsed_rows.append(cells)

                if parsed_rows:
                    num_cols = max(len(r) for r in parsed_rows)
                    tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(tbl, "CBD5E1", "6")

                    col_width = Inches(6.5 / max(num_cols, 1))

                    for r_idx, row_data in enumerate(parsed_rows):
                        for c_idx in range(num_cols):
                            cell = tbl.cell(r_idx, c_idx)
                            cell.width = col_width
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)

                            cell_text = row_data[c_idx] if c_idx < len(row_data) else ""

                            if r_idx == 0:
                                # Header Row
                                set_cell_background(cell, "1E3A8A")  # Deep Blue
                                add_formatted_text(p, cell_text, default_size=9.5, default_color=RGBColor(255, 255, 255))
                                for r in p.runs:
                                    r.bold = True
                            else:
                                # Data Row (zebra striping)
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F8FAFC")
                                else:
                                    set_cell_background(cell, "FFFFFF")
                                add_formatted_text(p, cell_text, default_size=9.0)

                    # Space after table
                    p_post = doc.add_paragraph()
                    p_post.paragraph_format.space_before = Pt(4)
                    p_post.paragraph_format.space_after = Pt(4)
            continue

        # 5. Blockquote (> ...)
        if line.startswith(">"):
            quote_lines = []
            while i < total_lines and lines[i].rstrip().startswith(">"):
                quote_lines.append(lines[i].rstrip()[1:].strip())
                i += 1

            quote_text = " ".join(quote_lines)

            # Create a 1x1 callout table with blue left border
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.5)
            set_cell_background(cell, "F0F9FF")  # Light sky tint
            set_callout_borders(tbl, "2563EB")  # Solid blue left bar
            set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, quote_text, default_size=10.0, default_color=RGBColor(30, 58, 138))
            for r in p.runs:
                r.italic = True

            p_post = doc.add_paragraph()
            p_post.paragraph_format.space_before = Pt(4)
            p_post.paragraph_format.space_after = Pt(4)
            continue

        # 6. Images (![alt](path))
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            alt_text = img_match.group(1)
            img_rel_path = img_match.group(2).strip()

            img_full_path = os.path.normpath(os.path.join(base_dir, img_rel_path))

            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                try:
                    run_img.add_picture(img_full_path, width=Inches(6.2))
                    print(f"    [+] Embedded image: {img_rel_path}")
                except Exception as e:
                    print(f"    [-] Failed to embed image {img_rel_path}: {e}")
                    run_img.add_text(f"[Image: {alt_text}]")
            else:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_img = p_img.add_run(f"[Image Missing: {img_rel_path}]")
                run_img.font.color.rgb = RGBColor(239, 68, 68)

            # Check if next line is a figure caption (*Figure: ...*)
            if i + 1 < total_lines and lines[i + 1].strip().startswith("*Figure") and lines[i + 1].strip().endswith("*"):
                i += 1
                caption_text = lines[i].strip()[1:-1]
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(12)
                r_cap = p_cap.add_run(caption_text)
                r_cap.font.name = "Calibri"
                r_cap.font.size = Pt(9.0)
                r_cap.italic = True
                r_cap.font.color.rgb = RGBColor(107, 114, 128)  # Gray-500

            i += 1
            continue

        # 7. Headings (#, ##, ###, ####, #####)
        if line.startswith("# "):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(title_text)
            run.font.name = "Calibri"
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy Blue
            i += 1
            continue

        if line.startswith("## "):
            h2_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h2_text)
            run.font.name = "Calibri"
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235)  # Royal Blue
            i += 1
            continue

        if line.startswith("### "):
            h3_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h3_text)
            run.font.name = "Calibri"
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(30, 64, 175)  # Navy
            i += 1
            continue

        if line.startswith("#### "):
            h4_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(h4_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11.5)
            run.bold = True
            run.font.color.rgb = RGBColor(55, 65, 81)  # Dark Gray
            i += 1
            continue

        if line.startswith("##### "):
            h5_text = line[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h5_text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = RGBColor(75, 85, 99)
            i += 1
            continue

        # 8. Unordered Lists (- ..., * ...)
        if re.match(r'^\s*[-*+]\s+', line):
            indent_level = (len(line) - len(line.lstrip())) // 2
            list_text = re.sub(r'^\s*[-*+]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            add_formatted_text(p, list_text, default_size=10.5)
            i += 1
            continue

        # 9. Ordered Lists (1. ..., 2. ...)
        if re.match(r'^\s*\d+\.\s+', line):
            indent_level = (len(line) - len(line.lstrip())) // 2
            list_text = re.sub(r'^\s*\d+\.\s+', '', line)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            add_formatted_text(p, list_text, default_size=10.5)
            i += 1
            continue

        # 10. Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, line, default_size=10.5)
        i += 1

    print(f"[*] Saving Word Document to: {docx_path}")
    doc.save(docx_path)
    print(f"[OK] Document successfully generated: {docx_path}")


if __name__ == "__main__":
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    md_file = os.path.join(base_dir, "MEDIUM_TUTORIAL_BLOG_POST.md")
    output_docx = os.path.join(base_dir, "MEDIUM_TUTORIAL_BLOG_POST.docx")
    friendly_docx = os.path.join(base_dir, "NexusLake_Medium_Tutorial_Article.docx")

    convert_markdown_to_docx(md_file, output_docx, base_dir)

    # Save a readable copy as well
    import shutil
    shutil.copy2(output_docx, friendly_docx)
    print(f"[OK] Friendly copy created: {friendly_docx}")
